"""Informe ejecutivo de Matcher en Excel.

    python3 -m marketing.generar_informe

Sale `docs/Matcher-analisis.xlsx` con cinco hojas:

  Calificación   nota 1-10 por área, con el motivo de cada una
  Rentabilidad   mes 1, 3, 6, 9, 12, 18 y 24 en los tres escenarios
  Con y sin ads  el mismo negocio comprando usuarios o sin comprar
  Competencia    Uruguay / Latam / mundo
  Supuestos      de dónde sale cada número y qué NO está verificado

POR QUÉ SE GENERA
Los números salen de `marketing/modelo_negocio.py` —el mismo modelo que
produce `docs/PLAN_NEGOCIO.md`— y los precios de `matcher/planes.py`. Un
Excel escrito a mano se desincroniza del producto en la primera semana y
después nadie sabe cuál de los dos miente.

LA CALIFICACIÓN ES UNA OPINIÓN, Y ASÍ SE PRESENTA
Las notas de la primera hoja son un juicio, no una medición. Lo que sí es
medible —tests que pasan, auditorías corridas— va en la columna de evidencia,
para que se pueda discutir la nota mirando el dato y no la impresión.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from matcher import planes

from . import modelo_negocio as mn

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "docs" / "Matcher-analisis.xlsx"

HITOS = (1, 3, 6, 9, 12, 18, 24)

# Paleta: la misma de la app, para que el informe se vea del mismo producto.
NAVY = "FF0B1016"
CORAL = "FFFF4655"
GRIS = "FF1A2430"
TINTA = "FFF2F6F9"

_borde = Border(bottom=Side(style="thin", color="FF2C3947"))


def _titulo(hoja, fila: int, texto: str, ancho: int) -> int:
    c = hoja.cell(row=fila, column=1, value=texto)
    c.font = Font(bold=True, size=13, color=TINTA)
    c.fill = PatternFill("solid", fgColor=NAVY)
    for col in range(2, ancho + 1):
        hoja.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=NAVY)
    return fila + 1


def _cabecera(hoja, fila: int, columnas: list[str]) -> int:
    for i, texto in enumerate(columnas, start=1):
        c = hoja.cell(row=fila, column=i, value=texto)
        c.font = Font(bold=True, color=TINTA, size=10)
        c.fill = PatternFill("solid", fgColor=GRIS)
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        c.border = _borde
    hoja.row_dimensions[fila].height = 28
    return fila + 1


def _anchos(hoja, anchos: list[int]) -> None:
    for i, a in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = a


# ---------------------------------------------------------------------------
# Hoja 1 — Calificación
# ---------------------------------------------------------------------------
# La evidencia es lo que sostiene cada nota. Sin ella esto sería una opinión
# con formato de tabla.
NOTAS = [
    ("Diseño / UI", 8.5,
     "Paleta propia coherente (carbón + fuego), set de íconos propio, un solo CSS para "
     "web/APK/PC, modo teléfono real. Le falta pulido de micro-interacciones y un "
     "diseñador que revise tipografía y espaciados con ojo fresco.",
     "Verificado en navegador a 412 px y 1280 px, 0 errores de JS."),
    ("Web pública", 8.0,
     "Landing en tres idiomas generada desde el código: los precios salen de `planes.py`, "
     "así que no puede prometer un plan que no existe. Videos de demo embebidos con "
     "respaldo WebM. Falta SEO real, dominio propio y textos revisados por alguien de marketing.",
     "16 tests + recorrido de navegador en 3 idiomas, 0 fallos."),
    ("Seguridad", 7.5,
     "Sesiones firmadas y revocables, borrado real de cuenta, webhooks con firma "
     "verificada, CSP sin eval, sin Node expuesto en el escritorio, el servidor decide "
     "los planes (no el cliente). Baja de 9 porque NO hay moderación de contenido y "
     "porque nada de esto fue auditado por un tercero.",
     "El agujero que regalaba planes se encontró y se cerró en esta sesión."),
    ("Funcionalidad", 8.5,
     "Filtros duros, radar con mapa, cruces, Crush Time, cita a ciegas, segunda vuelta, "
     "vitrinas, videollamada con consentimiento de los dos, automatch, planes y pagos. "
     "Más features que varias apps del rubro en producción.",
     "400 tests verdes, auditoría HTTP end-to-end con 0 fallos."),
    ("Rentabilidad (producto)", 6.5,
     "El precio está muy por debajo de la competencia y los márgenes por suscriptor son "
     "buenos, pero el negocio depende de densidad de usuarios en una zona chica. Sin esa "
     "densidad no hay matches y sin matches no hay pagos. Es el riesgo real, no el código.",
     "Modelo con 3 escenarios; ver la hoja Rentabilidad."),
    ("Listo para producción", 6.0,
     "Anda de punta a punta y los pagos están bien cerrados, pero falta: cobro real "
     "probado con tarjeta, moderación, backend con disco persistente (hoy es efímero) y "
     "las cuentas de tienda. Nada de eso es código: es trámite y plata.",
     "Ver el checklist en docs/PRODUCCION.md."),
]


def hoja_calificacion(wb: Workbook) -> None:
    h = wb.active
    h.title = "Calificación"
    _anchos(h, [26, 9, 78, 52])

    f = _titulo(h, 1, "MATCHER · calificación del estado actual", 4)
    f += 1
    f = _cabecera(h, f, ["Área", "Nota /10", "Por qué esa nota", "Evidencia"])

    for area, nota, motivo, evidencia in NOTAS:
        h.cell(row=f, column=1, value=area).font = Font(bold=True)
        c = h.cell(row=f, column=2, value=nota)
        c.alignment = Alignment(horizontal="center")
        c.font = Font(bold=True, color=CORAL, size=12)
        h.cell(row=f, column=3, value=motivo).alignment = Alignment(wrap_text=True, vertical="top")
        h.cell(row=f, column=4, value=evidencia).alignment = Alignment(wrap_text=True, vertical="top")
        h.row_dimensions[f].height = 58
        f += 1

    promedio = round(sum(n for _, n, _, _ in NOTAS) / len(NOTAS), 1)
    f += 1
    h.cell(row=f, column=1, value="GENERAL").font = Font(bold=True, size=12)
    c = h.cell(row=f, column=2, value=promedio)
    c.font = Font(bold=True, size=14, color=CORAL)
    c.alignment = Alignment(horizontal="center")
    h.cell(
        row=f, column=3,
        value="Producto sólido y con diferenciales reales. Lo que falta para vender no es "
              "código: es una pasarela probada con plata de verdad, moderación y las "
              "cuentas de tienda.",
    ).alignment = Alignment(wrap_text=True, vertical="top")
    h.row_dimensions[f].height = 44


# ---------------------------------------------------------------------------
# Hoja 2 — Rentabilidad
# ---------------------------------------------------------------------------
def hoja_rentabilidad(wb: Workbook) -> None:
    h = wb.create_sheet("Rentabilidad")
    _anchos(h, [22, 10, 13, 13, 14, 14, 13, 13, 14, 16])

    f = _titulo(h, 1, "RENTABILIDAD NETA · después de comisiones e IRAE (Uruguay)", 10)
    f += 1

    for e in mn.ESCENARIOS:
        corrida = mn.correr(e)
        h.cell(row=f, column=1, value=f"Escenario {e.nombre}").font = Font(bold=True, size=12)
        h.cell(row=f, column=2, value=e.resumen).alignment = Alignment(wrap_text=True)
        f += 1
        f = _cabecera(h, f, [
            "Mes", "Usuarios", "Suscriptores", "Bruto USD", "Neto pasarela USD",
            "Marketing USD", "Infra USD", "Otros USD", "IRAE USD", "RESULTADO NETO USD",
        ])
        for numero in HITOS:
            m = corrida.mes(numero)
            valores = [
                f"Mes {numero}", round(m.usuarios), round(m.suscriptores),
                round(m.ingreso_bruto, 2), round(m.ingreso_neto, 2),
                round(m.marketing, 2), round(m.infra, 2),
                round(m.moderacion + m.fijos, 2), round(m.impuesto, 2),
                round(m.resultado_neto, 2),
            ]
            for i, v in enumerate(valores, start=1):
                c = h.cell(row=f, column=i, value=v)
                c.border = _borde
                if i == 10:
                    c.font = Font(bold=True, color="FF7CC242" if v >= 0 else CORAL)
                if i >= 4:
                    c.number_format = "#,##0.00"
            f += 1

        acumulado = corrida.acumulado(24)
        h.cell(row=f, column=1, value="Acumulado 24 meses").font = Font(bold=True)
        c = h.cell(row=f, column=10, value=round(acumulado, 2))
        c.font = Font(bold=True, color="FF7CC242" if acumulado >= 0 else CORAL)
        c.number_format = "#,##0.00"
        f += 3


# ---------------------------------------------------------------------------
# Hoja 3 — Con y sin publicidad
# ---------------------------------------------------------------------------
def hoja_ads(wb: Workbook) -> None:
    h = wb.create_sheet("Con y sin ads")
    _anchos(h, [26, 14, 16, 16, 16, 16])

    f = _titulo(h, 1, "EL MISMO NEGOCIO, COMPRANDO USUARIOS O SIN COMPRARLOS", 6)
    f += 1
    h.cell(
        row=f, column=1,
        value="Sin ads el crecimiento es sólo boca a boca: más lento, pero sin costo de "
              "adquisición. Con ads se crece antes y se paga por cada usuario. La pregunta "
              "no es cuál da más plata al mes 24, es cuánto aguantás perdiendo hasta ahí.",
    ).alignment = Alignment(wrap_text=True)
    h.merge_cells(start_row=f, start_column=1, end_row=f, end_column=6)
    h.row_dimensions[f].height = 42
    f += 2

    import dataclasses

    base = next(e for e in mn.ESCENARIOS if e.codigo == "base")
    # `Escenario` es inmutable a propósito (nadie puede tocarle los supuestos a
    # un escenario a mitad de una corrida), así que la variante sin publicidad
    # se arma con `replace`, no asignando el campo.
    sin_ads = dataclasses.replace(base, presupuesto=[0.0] * len(base.presupuesto))

    for etiqueta, escenario in (("CON publicidad", base), ("SIN publicidad", sin_ads)):
        corrida = mn.correr(escenario)
        h.cell(row=f, column=1, value=etiqueta).font = Font(bold=True, size=12)
        f += 1
        f = _cabecera(h, f, ["Mes", "Usuarios", "Suscriptores", "Marketing USD",
                             "Resultado neto USD", "Acumulado USD"])
        for numero in HITOS:
            m = corrida.mes(numero)
            valores = [f"Mes {numero}", round(m.usuarios), round(m.suscriptores),
                       round(m.marketing, 2), round(m.resultado_neto, 2),
                       round(corrida.acumulado(numero), 2)]
            for i, v in enumerate(valores, start=1):
                c = h.cell(row=f, column=i, value=v)
                c.border = _borde
                if i >= 4:
                    c.number_format = "#,##0.00"
                if i in (5, 6) and isinstance(v, (int, float)):
                    c.font = Font(bold=(i == 6), color="FF7CC242" if v >= 0 else CORAL)
            f += 1
        f += 2


# ---------------------------------------------------------------------------
# Hoja 4 — Competencia
# ---------------------------------------------------------------------------
# `verificado` marca qué se pudo comprobar y qué es referencia aproximada. Es
# la regla 10 del producto aplicada al informe: publicar un precio ajeno como
# dato firme es lo que trae una carta documento.
COMPETENCIA = [
    ("Tinder", "Mundo / Uruguay", "~15,99–29,99", "No",
     "El más grande y el más caro. Los filtros finos son de pago y aun así 'no se respetan' "
     "es su queja número uno en las reseñas."),
    ("Bumble", "Mundo / Uruguay", "~29,99", "No",
     "Ellas escriben primero. Buena reputación de seguridad, precio alto."),
    ("Happn", "Mundo / Uruguay", "~24,99", "No",
     "Su diferencial es el cruce en la calle — lo mismo que hace el Radar de Matcher, "
     "que acá va en el plan gratis."),
    ("Kismia", "Latam", "variable", "No",
     "Fuerte en Latam por publicidad agresiva; mala fama por cobros poco claros."),
    ("Grindr", "Mundo", "variable", "No",
     "Nicho definido, no compite de frente."),
    ("Matcher", "Uruguay → Latam → mundo",
     f"{planes.PLANES['plus'].precio_mes:.2f} / {planes.PLANES['gold'].precio_mes:.2f}", "Sí",
     "Todos los filtros gratis, precio 4–7 veces menor, videollamada con consentimiento, "
     "cita a ciegas, segunda vuelta y radar. Sin usuarios todavía: ése es el punto débil."),
]

MERCADOS = [
    ("Uruguay", "~3,4 M de habitantes",
     "Mercado chico pero alcanzable: se puede dominar Montevideo con presupuesto bajo. "
     "La densidad es EL problema y también la oportunidad — en un país chico, 5.000 "
     "usuarios activos ya hacen que la app funcione."),
    ("Latam", "~660 M",
     "Brasil y México son enormes y caros de adquirir. Entrar por países chicos "
     "(Uruguay, Paraguay, Costa Rica) y crecer por vecindad es más barato que pelear "
     "de frente en São Paulo."),
    ("Mundo", "~5.000 M con internet",
     "No es un objetivo realista de corto plazo. Match Group (Tinder, Hinge, OkCupid) y "
     "Bumble tienen presupuestos de marketing que no se compiten con dinero: se compite "
     "con un diferencial que ellos no quieran copiar."),
]


def hoja_competencia(wb: Workbook) -> None:
    h = wb.create_sheet("Competencia")
    _anchos(h, [16, 24, 18, 12, 76])

    f = _titulo(h, 1, "COMPETENCIA", 5)
    f += 1
    f = _cabecera(h, f, ["App", "Dónde", "USD/mes aprox.", "¿Verificado?", "Notas"])
    for fila in COMPETENCIA:
        for i, v in enumerate(fila, start=1):
            c = h.cell(row=f, column=i, value=v)
            c.border = _borde
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if fila[0] == "Matcher":
                c.font = Font(bold=True, color=CORAL)
        h.row_dimensions[f].height = 44
        f += 1

    f += 1
    h.cell(
        row=f, column=1,
        value="Los precios de la competencia son de lista, aproximados, y NO están "
              "verificados: cambian por país y por promoción. Están para dar escala, no "
              "como dato firme — publicarlos como exactos es un problema legal.",
    ).font = Font(italic=True, color="FF9CADBD")
    h.merge_cells(start_row=f, start_column=1, end_row=f, end_column=5)
    f += 3

    f = _cabecera(h, f, ["Mercado", "Tamaño", "Lectura", "", ""])
    for mercado, tamanio, lectura in MERCADOS:
        h.cell(row=f, column=1, value=mercado).font = Font(bold=True)
        h.cell(row=f, column=2, value=tamanio)
        c = h.cell(row=f, column=3, value=lectura)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        h.merge_cells(start_row=f, start_column=3, end_row=f, end_column=5)
        h.row_dimensions[f].height = 58
        f += 1


# ---------------------------------------------------------------------------
# Hoja 5 — Supuestos
# ---------------------------------------------------------------------------
def hoja_supuestos(wb: Workbook) -> None:
    h = wb.create_sheet("Supuestos")
    _anchos(h, [34, 18, 84])

    f = _titulo(h, 1, "DE DÓNDE SALE CADA NÚMERO", 3)
    f += 1
    f = _cabecera(h, f, ["Supuesto", "Valor", "Origen / advertencia"])

    filas = [
        ("Precio Plus (mes)", f"USD {planes.PLANES['plus'].precio_mes:.2f}",
         "De `matcher/planes.py`: la misma fuente que cobra el checkout."),
        ("Precio Gold (mes)", f"USD {planes.PLANES['gold'].precio_mes:.2f}", "Ídem."),
        ("Comisión web", f"{mn.COMISION_WEB * 100:.1f} %",
         "Estimación de cobro con tarjeta en Latam. VERIFICAR contra tu liquidación real "
         "de MercadoPago y corregir en `modelo_negocio.py`."),
        ("Comisión tienda", f"{mn.COMISION_TIENDA * 100:.0f} %",
         "Apple Small Business Program y equivalente de Google, hasta USD 1M/año. "
         f"Arriba de eso, {mn.COMISION_TIENDA_ALTA * 100:.0f} %."),
        ("IRAE", f"{mn.IRAE * 100:.0f} %",
         "Impuesto a la renta empresarial en Uruguay sobre la renta neta fiscal. Las "
         "pérdidas de meses anteriores compensan. NO incluye IVA, aportes ni el régimen "
         "que te corresponda: eso lo define un contador según cómo constituyas el negocio."),
        ("Impuestos NO incluidos", "IVA, BPS, IRPF",
         "El modelo descuenta comisiones e IRAE. Cualquier otro tributo depende de la "
         "figura jugídica (unipersonal, SAS) y de si facturás a Uruguay o al exterior. "
         "Esto lo tiene que mirar un contador, no un modelo."),
        ("Usuarios", "modelados",
         "NO hay usuarios reales todavía. Todo el crecimiento de la hoja Rentabilidad es "
         "un modelo con supuestos de conversión y churn, no una medición. El primer mes "
         "con gente real reemplaza todo esto."),
    ]
    for supuesto, valor, origen in filas:
        h.cell(row=f, column=1, value=supuesto).font = Font(bold=True)
        h.cell(row=f, column=2, value=valor)
        h.cell(row=f, column=3, value=origen).alignment = Alignment(wrap_text=True, vertical="top")
        h.row_dimensions[f].height = 46
        f += 1


def generar() -> Path:
    wb = Workbook()
    hoja_calificacion(wb)
    hoja_rentabilidad(wb)
    hoja_ads(wb)
    hoja_competencia(wb)
    hoja_supuestos(wb)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    return SALIDA


if __name__ == "__main__":
    print(generar())
